import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
import os
from datetime import datetime

# --- إعداد ملف Excel ---
excel_file = "supermarket_data.xlsx"

# إنشاء الملف إذا لم يكن موجود
if not os.path.exists(excel_file):
    wb = Workbook()
    ws_customers = wb.active
    ws_customers.title = "Customers"
    ws_customers.append(["Customer_ID", "Customer_Name", "Phone"])

    ws_products = wb.create_sheet("Products")
    ws_products.append(["Product_ID", "Product_Name", "Price", "Quantity"])

    ws_sales = wb.create_sheet("Sales")
    ws_sales.append(["Sale_ID", "Product_ID", "Quantity", "Sale_Date"])

    ws_invoices = wb.create_sheet("Invoices")
    ws_invoices.append(["Invoice_ID", "Product_ID", "Product_Name", "Price", "Quantity", "Total", "Date"])

    wb.save(excel_file)

# تحميل الملف
wb = load_workbook(excel_file)

# التأكد من وجود ورقة "Invoices" وإذا لم تكن موجودة، يتم إنشاؤها
if "Invoices" not in wb.sheetnames:
    ws_invoices = wb.create_sheet("Invoices")
    ws_invoices.append(["Invoice_ID", "Product_ID", "Product_Name", "Price", "Quantity", "Total", "Date"])
else:
    ws_invoices = wb["Invoices"]

ws_customers = wb["Customers"]
ws_products = wb["Products"]
ws_sales = wb["Sales"]

# --- إعداد الواجهة ---
root = tk.Tk()
root.title("نظام إدارة السوبر ماركت")
root.geometry("900x600")
root.configure(bg="#f5f5f5")

style = ttk.Style()
style.theme_use("clam")
style.configure("TButton",
                font=("Arial", 12, "bold"),
                padding=10,
                background="#4CAF50",
                foreground="white",
                relief="flat",
                width=20)
style.configure("TLabel", font=("Arial", 14), background="#f5f5f5", foreground="#333")
style.configure("TFrame", background="#f5f5f5")

main_frame = ttk.Frame(root, padding=20)
main_frame.pack(expand=True, fill=tk.BOTH)

# --- دوال ---
def get_product_details(product_id):
    for row in ws_products.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == str(product_id):
            return row  # [ID, Name, Price, Quantity]
    return None

def update_product_quantity(product_id, quantity_sold):
    for row in ws_products.iter_rows(min_row=2, values_only=False):
        if str(row[0].value) == str(product_id):
            new_quantity = row[3].value - quantity_sold
            row[3].value = new_quantity
            wb.save(excel_file)
            return True
    return False

def add_invoice():
    win = tk.Toplevel(root)
    win.title("إضافة فاتورة")
    win.geometry("400x350")
    win.configure(bg="#e0f7fa")

    ttk.Label(win, text="رقم المنتج:", background="#e0f7fa").pack(pady=5)
    product_entry = ttk.Entry(win)
    product_entry.pack()

    ttk.Label(win, text="الكمية:", background="#e0f7fa").pack(pady=5)
    quantity_entry = ttk.Entry(win)
    quantity_entry.pack()

    ttk.Label(win, text="السعر الإجمالي:", background="#e0f7fa").pack(pady=5)
    total_label = ttk.Label(win, text="0.00", background="#e0f7fa")
    total_label.pack()

    # دالة لحساب السعر الإجمالي عند إدخال الكمية
    def calculate_total():
        pid = product_entry.get()
        qty = quantity_entry.get()

        if pid and qty:
            product = get_product_details(pid)
            if product:
                try:
                    qty = int(qty)
                    price = float(product[2])
                    total = price * qty
                    total_label.config(text=f"{total:.2f}")  # تحديث السعر الإجمالي
                except ValueError:
                    messagebox.showwarning("تحذير", "الكمية يجب أن تكون رقمًا صحيحًا")
            else:
                messagebox.showerror("خطأ", "المنتج غير موجود")
        else:
            total_label.config(text="0.00")

    # تحديث السعر الإجمالي عند إدخال الكمية أو رقم المنتج
    product_entry.bind("<KeyRelease>", lambda event: calculate_total())
    quantity_entry.bind("<KeyRelease>", lambda event: calculate_total())

    # دالة لحفظ الفاتورة كـ PDF
    def save_invoice_pdf(data):
        if not os.path.exists("invoices"):
            os.makedirs("invoices")
        
        now = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"invoices/invoice_{now}.pdf"
        pdf = canvas.Canvas(filename, pagesize=A4)
        width, height = A4

        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(200, height - 50, "فاتورة البيع")
        pdf.setFont("Helvetica", 12)

        y = height - 100
        for row in data:
            line = f"رقم المنتج: {row[0]} - الاسم: {row[1]} - السعر: {row[2]} - الكمية: {row[3]} - الإجمالي: {row[4]}"
            pdf.drawString(50, y, line)
            y -= 25
        pdf.save()

    def save():
        pid = product_entry.get()
        qty = quantity_entry.get()
        date_now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if not pid or not qty:
            messagebox.showwarning("تحذير", "أدخل جميع البيانات")
            return

        product = get_product_details(pid)
        if not product:
            messagebox.showerror("خطأ", "المنتج غير موجود")
            return

        try:
            qty = int(qty)
            price = float(product[2])
            total = price * qty

            # إضافة الفاتورة في ورقة "Invoices"
            invoice_id = ws_invoices.max_row
            ws_invoices.append([invoice_id, pid, product[1], price, qty, total, date_now])
            ws_sales.append([invoice_id, pid, qty, date_now])
            wb.save(excel_file)

            # تحديث المخزون بعد بيع المنتج
            if not update_product_quantity(pid, qty):
                messagebox.showerror("خطأ", "لم يتم تحديث المخزون")
                return

            # حفظ الفاتورة PDF
            save_invoice_pdf([[pid, product[1], price, qty, total]])

            messagebox.showinfo("تم", f"تمت إضافة الفاتورة وحفظها كـ PDF\nالإجمالي: {total} جنيه")
            win.destroy()
        except ValueError:
            messagebox.showerror("خطأ", "الكمية يجب أن تكون رقمًا صحيحًا")
        except Exception as e:
            messagebox.showerror("خطأ", str(e))

    ttk.Button(win, text="إضافة و حفظ الفاتورة PDF", command=save).pack(pady=20)

# تحسين الأزرار والنصوص
def add_customer():
    win = tk.Toplevel(root)
    win.title("إضافة عميل")
    win.geometry("350x200")
    win.configure(bg="#e1bee7")

    ttk.Label(win, text="اسم العميل:", background="#e1bee7").pack(pady=5)
    name_entry = ttk.Entry(win)
    name_entry.pack()

    ttk.Label(win, text="رقم الهاتف:", background="#e1bee7").pack(pady=5)
    phone_entry = ttk.Entry(win)
    phone_entry.pack()

    def save():
        name, phone = name_entry.get(), phone_entry.get()
        if name and phone:
            new_id = ws_customers.max_row
            ws_customers.append([new_id, name, phone])
            wb.save(excel_file)
            messagebox.showinfo("تم", "تمت إضافة العميل")
            win.destroy()
        else:
            messagebox.showwarning("تحذير", "أدخل جميع البيانات")

    ttk.Button(win, text="حفظ", command=save).pack(pady=10)

def add_product():
    win = tk.Toplevel(root)
    win.title("إضافة منتج")
    win.geometry("350x250")
    win.configure(bg="#e3f2fd")

    ttk.Label(win, text="اسم المنتج:", background="#e3f2fd").pack(pady=5)
    name_entry = ttk.Entry(win)
    name_entry.pack()

    ttk.Label(win, text="السعر:", background="#e3f2fd").pack(pady=5)
    price_entry = ttk.Entry(win)
    price_entry.pack()

    ttk.Label(win, text="الكمية:", background="#e3f2fd").pack(pady=5)
    quantity_entry = ttk.Entry(win)
    quantity_entry.pack()

    def save():
        name, price, quantity = name_entry.get(), price_entry.get(), quantity_entry.get()
        if name and price and quantity:
            new_id = ws_products.max_row
            ws_products.append([new_id, name, price, quantity])
            wb.save(excel_file)
            messagebox.showinfo("تم", "تمت إضافة المنتج")
            win.destroy()
        else:
            messagebox.showwarning("تحذير", "أدخل جميع البيانات")

    ttk.Button(win, text="حفظ", command=save).pack(pady=10)

def show_table(title, columns, sheet):
    win = tk.Toplevel(root)
    win.title(title)
    win.geometry("700x400")

    tree = ttk.Treeview(win, columns=columns, show="headings")
    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=150)
    tree.pack(fill=tk.BOTH, expand=True)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)

def view_inventory():
    show_table("المخزن", ("ID", "المنتج", "السعر", "الكمية"), ws_products)

def view_sales():
    show_table("المبيعات", ("ID", "رقم المنتج", "الكمية", "التاريخ"), ws_sales)

# --- الواجهة الرئيسية ---
ttk.Label(main_frame, text="نظام إدارة السوبر ماركت", font=("Arial", 18, "bold"), foreground="#0277bd").pack(pady=20)

ttk.Button(main_frame, text="إضافة عميل", command=add_customer).pack(fill='x', pady=10)
ttk.Button(main_frame, text="إضافة منتج", command=add_product).pack(fill='x', pady=10)
ttk.Button(main_frame, text="إضافة فاتورة", command=add_invoice).pack(fill='x', pady=10)
ttk.Button(main_frame, text="عرض المخزن", command=view_inventory).pack(fill='x', pady=10)
ttk.Button(main_frame, text="عرض المبيعات", command=view_sales).pack(fill='x', pady=10)

root.mainloop()
